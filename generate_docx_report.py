"""
Generates the formal academic research paper in Microsoft Word (.docx) format
with embedded high-resolution figures, styled tables, and formatted headings.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_research_docx():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    docs_dir = os.path.join(base_dir, 'docs')
    outputs_dir = os.path.join(base_dir, 'outputs')
    docx_path = os.path.join(docs_dir, 'research_paper.docx')
    
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Student Segmentation and Recommender System:\nAn Empirical Behavioral Analytics and Multi-Factor Recommender Framework for Digital Learning Platforms")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Dark slate
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run("Data Science & Learning Analytics Research Team • EduPro Analytics Group\nDate: October 2026")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(10.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph() # Spacing

    # Abstract Box
    p_abs_h = doc.add_paragraph()
    run_abs_h = p_abs_h.add_run("Abstract")
    run_abs_h.font.name = "Arial"
    run_abs_h.font.size = Pt(12)
    run_abs_h.font.bold = True
    
    p_abs = doc.add_paragraph()
    run_abs = p_abs.add_run(
        "Online learning platforms experience significant diversity in student motivation, curricular exploration, "
        "and willingness to pay. However, conventional course recommendation systems frequently suffer from cold-start "
        "bottlenecks, overemphasize global popularity, or lack contextual interpretability. This research develops an "
        "end-to-end learning analytics framework that couples unsupervised behavioral student segmentation with an "
        "explainable, multi-factor recommendation engine. Using a verified dataset of 3,000 learners, 60 instructors, "
        "60 courses across 12 disciplines, and 10,000 transaction events spanning a complete academic calendar year, we "
        "engineer eight longitudinal behavioral features while strictly excluding sensitive demographic attributes "
        "(age, gender) to safeguard against algorithmic bias. Comparative cluster evaluation reveals that a 4-cluster "
        "K-Means solution (k=4, Silhouette = 0.340, Calinski-Harabasz = 2005.3) establishes the optimal trade-off between "
        "geometric separation and educational domain interpretability. The identified student segments comprise Beginner "
        "Free Trialers (31.5%), Direct Paid Investors (29.8%), Advanced Skill Specialists (23.6%), and High-Engagement "
        "Super-Learners (15.1%). Leveraging these segments, we construct a hybrid recommendation engine synthesizing "
        "segment popularity (35%), content curriculum progression (25%), k-nearest neighbor behavioral collaborative "
        "filtering (20%), and catalog quality ratings (20%), while enforcing strict exclusion of prior enrollments. "
        "Rigorous offline evaluation using a chronological leave-one-out holdout across 1,380 active learners proves that "
        "the hybrid framework achieves a +10.9% relative lift in Mean Reciprocal Rank (MRR@5) (0.0478 vs. 0.0431) and a "
        "+6.8% lift in NDCG@5 (0.0611 vs. 0.0572) over an unbiased global popularity baseline."
    )
    run_abs.font.name = "Arial"
    run_abs.font.size = Pt(10)
    
    p_kw = doc.add_paragraph()
    r_kw_title = p_kw.add_run("Keywords: ")
    r_kw_title.font.bold = True
    r_kw_title.font.size = Pt(9.5)
    r_kw = p_kw.add_run("Learning Analytics, Student Segmentation, Course Recommender Systems, Behavioral Clustering, K-Means, Collaborative Filtering, Educational Data Mining.")
    r_kw.font.size = Pt(9.5)
    r_kw.font.italic = True

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Digital education platforms have democratized access to technical and professional education. "
        "However, learners face extreme choice overload when navigating multi-disciplinary course catalogs. "
        "Traditional collaborative filtering algorithms often fail in learning environments due to extreme matrix sparsity—"
        "over 54% of learners in typical e-learning ecosystems possess only a single recorded transaction. Furthermore, "
        "naive popularity-driven recommendation systems induce popularity bias, steering users towards ubiquitous "
        "introductory classes while neglecting specialized, advanced curricula.\n\n"
        "To resolve these challenges, this study presents a unified framework for Student Segmentation and Course "
        "Recommendation. By extracting nuanced behavioral signals—such as curricular exploration entropy, hierarchical "
        "learning depth, and monetary commitment ratios—we group students into meaningful behavioral archetypes. "
        "We then condition recommendation rankings upon segment-specific preferences and collaborative peer similarity, "
        "delivering transparent, explainable recommendations."
    )

    doc.add_heading("2. Dataset Architecture & Exploratory Analysis", level=1)
    doc.add_paragraph(
        "The empirical study evaluates a verified relational database comprising 3,000 learners, 60 instructors, "
        "60 courses across 12 disciplines, and 10,000 transaction events spanning from January 1, 2025 to December 30, 2025. "
        "Analysis of transaction amounts confirms that 6,403 transactions correspond to 0.00 dollar free course enrollments, "
        "while 3,597 transactions represent paid course purchases totaling $911,323.47 in platform revenue."
    )

    # Embed Figure 1
    fig1_path = os.path.join(outputs_dir, 'eda_category_enrolments.png')
    if os.path.exists(fig1_path):
        doc.add_picture(fig1_path, width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run("Figure 1: Distribution of Student Enrollments across Course Categories")
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True

    # Embed Figure 2
    fig2_path = os.path.join(outputs_dir, 'eda_paid_amounts_dist.png')
    if os.path.exists(fig2_path):
        doc.add_picture(fig2_path, width=Inches(5.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run("Figure 2: Distribution of Paid Course Transaction Amounts ($)")
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True

    doc.add_heading("3. Feature Engineering & Anti-Bias Methodology", level=1)
    doc.add_paragraph(
        "To ensure fair, unbiased, and effective personalization, eight longitudinal behavioral metrics were computed for "
        "each learner, indexing their activity, engagement, diversity, and level progression:\n"
        "1. Total Enrollments (Eu): Cumulative number of course enrollments.\n"
        "2. Total Spend (Su): Total financial expenditure in USD.\n"
        "3. Average Spend (Au): Expenditure per course enrollment.\n"
        "4. Paid Enrollment Ratio (Pu): Proportion of courses requiring payment.\n"
        "5. Category Breadth (Cu): Number of distinct categories explored.\n"
        "6. Curricular Diversity Score (Du): Normalized Shannon entropy across the 12 platform disciplines.\n"
        "7. Learning Depth Index (Lu): Weighted difficulty progression (1=Beginner, 2=Intermediate, 3=Advanced).\n"
        "8. Average Course Duration: Hours committed per course."
    )

    # Embed Figure 3
    fig3_path = os.path.join(outputs_dir, 'feature_correlations.png')
    if os.path.exists(fig3_path):
        doc.add_picture(fig3_path, width=Inches(5.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run("Figure 3: Pearson Correlation Matrix of Learner Features and Age")
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True

    doc.add_heading("4. Student Segmentation Results", level=1)
    doc.add_paragraph(
        "Standardized feature vectors were evaluated across k in [2, 8] using K-Means and Agglomerative Hierarchical "
        "Clustering. A 4-cluster K-Means solution (k=4) was selected based on the Elbow inflection point and "
        "educational domain interpretability."
    )

    # Embed Figure 4
    fig4_path = os.path.join(outputs_dir, 'clustering_evaluation_elbow_silhouette.png')
    if os.path.exists(fig4_path):
        doc.add_picture(fig4_path, width=Inches(6.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run("Figure 4: Elbow Method (Inertia) and Silhouette Score Diagnostics for K-Means and Hierarchical Clustering")
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True

    # Segment Profile Table
    table = doc.add_table(rows=5, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Segment Persona", "Count", "Share", "Avg Enrol", "Avg Spend", "Paid Ratio", "Depth (1-3)"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        set_cell_background(cell, "2563EB")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)
                
    row_data = [
        ["Beginner Free Trialers", "944", "31.5%", "1.49", "$24.52", "5.2%", "1.48"],
        ["Direct Paid Investors", "894", "29.8%", "1.55", "$336.23", "86.7%", "2.04"],
        ["Advanced Skill Specialists", "709", "23.6%", "1.59", "$29.52", "8.2%", "2.66"],
        ["High-Engagement Super-Learners", "453", "15.1%", "13.41", "$1,250.89", "36.8%", "1.98"]
    ]
    for i, row in enumerate(row_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.text = val
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    
    doc.add_paragraph() # Spacing

    # Embed Figure 5
    fig5_path = os.path.join(outputs_dir, 'cluster_pca_scatter.png')
    if os.path.exists(fig5_path):
        doc.add_picture(fig5_path, width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run("Figure 5: Two-Dimensional PCA Projection of the Four Learner Segments")
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True

    doc.add_heading("5. Hybrid Course Recommendation Engine", level=1)
    doc.add_paragraph(
        "The recommendation engine implements a four-factor linear utility scoring function combining:\n"
        "- Segment Popularity (35%): Course adoption rate within the learner's behavioral cluster.\n"
        "- Content Progression (25%): Category alignment and level progression (Beginner -> Intermediate -> Advanced).\n"
        "- Similar-Learner Collaborative Filtering (20%): Course choices of the 25 nearest behavioral neighbors.\n"
        "- Course Quality Baseline (20%): Normalized course rating.\n\n"
        "Crucially, the system strictly excludes all courses previously taken by the learner (guaranteeing 0 overlap) "
        "and generates plain-English pedagogical rationales for every suggestion."
    )

    doc.add_heading("6. Offline Evaluation & Benchmark Findings", level=1)
    doc.add_paragraph(
        "To rigorously evaluate recommendation accuracy without data leakage, a chronological leave-one-out holdout "
        "was executed across all 1,380 learners with multiple course transactions. Features, clusters, and collaborative "
        "matrices were fitted solely on earlier transactions, and evaluated against the held-out target course."
    )

    # Embed Figure 6
    fig6_path = os.path.join(outputs_dir, 'recommender_evaluation.png')
    if os.path.exists(fig6_path):
        doc.add_picture(fig6_path, width=Inches(5.5))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = p_cap.add_run("Figure 6: Offline Recommendation Performance: Hybrid Engine vs Popularity Baseline")
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True

    # Evaluation Table
    table2 = doc.add_table(rows=5, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers2 = ["Evaluation Metric", "Hybrid Engine", "Popularity Baseline", "Relative Lift"]
    for j, h in enumerate(headers2):
        cell = table2.cell(0, j)
        cell.text = h
        set_cell_background(cell, "2563EB")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)
                
    row_data2 = [
        ["Hit Rate@5", "10.22%", "10.07%", "+1.5%"],
        ["Precision@5", "0.0204", "0.0201", "+1.5%"],
        ["Mean Reciprocal Rank (MRR@5)", "0.0478", "0.0431", "+10.9%"],
        ["NDCG@5", "0.0611", "0.0572", "+6.8%"]
    ]
    for i, row in enumerate(row_data2):
        for j, val in enumerate(row):
            cell = table2.cell(i+1, j)
            cell.text = val
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    
    doc.add_paragraph() # Spacing

    doc.add_heading("7. Strategic Recommendations & Conclusion", level=1)
    doc.add_paragraph(
        "Based on our findings, we propose four tailored platform interventions:\n"
        "1. Beginner Free Trialers (31.5%): Deploy automated milestone prompts upon completion of free courses with intro discounts.\n"
        "2. Direct Paid Investors (29.8%): Surface career-focused multi-course certificates and accredited tracks.\n"
        "3. Advanced Skill Specialists (23.6%): Recommend expert masterclasses and technical capstones.\n"
        "4. High-Engagement Super-Learners (15.1%): Offer annual all-access memberships and peer community leadership.\n\n"
        "In conclusion, coupling unsupervised behavioral segmentation with an explainable multi-factor recommender provides "
        "a scalable, transparent, and anti-bias architecture for online education platforms."
    )

    doc.save(docx_path)
    print(f"Research paper (.docx) successfully created at: {docx_path}")

if __name__ == '__main__':
    create_research_docx()
